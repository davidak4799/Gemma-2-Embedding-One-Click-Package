# -*- coding: utf-8 -*-
"""
Google EmbeddingGemma 2 本地多模態搜索
文字 / 圖片 / 影片 / 音頻 → 同一個 768 維向量空間 → 互相搜索
"""
import os, sys, json, time, hashlib, subprocess, tempfile, platform
from pathlib import Path

import numpy as np
import streamlit as st

ROOT = Path(__file__).resolve().parent
MODEL_DIR = os.environ.get("EG2_MODEL_DIR", str(ROOT / "models" / "embeddinggemma-2"))
INDEX_DIR = Path(os.environ.get("EG2_INDEX_DIR", str(ROOT / "index")))
INDEX_DIR.mkdir(parents=True, exist_ok=True)

# ---------------- 檔案類型 ----------------
TEXT_EXT = {".txt", ".md", ".csv", ".json", ".log", ".ini", ".yaml", ".yml", ".xml", ".html", ".htm",
            ".py", ".js", ".ts", ".java", ".c", ".cpp", ".h", ".cs", ".go", ".rs", ".sql", ".bat",
            ".ps1", ".sh", ".css", ".tex", ".srt", ".docx", ".pdf", ".pptx", ".xlsx"}
IMAGE_EXT = {".jpg", ".jpeg", ".png", ".webp", ".bmp", ".gif", ".tif", ".tiff"}
AUDIO_EXT = {".mp3", ".wav", ".flac", ".m4a", ".aac", ".ogg", ".wma", ".opus"}
VIDEO_EXT = {".mp4", ".mkv", ".mov", ".avi", ".webm", ".wmv", ".flv", ".m4v", ".3gp"}
MODALITY = {**{e: "text" for e in TEXT_EXT}, **{e: "image" for e in IMAGE_EXT},
            **{e: "audio" for e in AUDIO_EXT}, **{e: "video" for e in VIDEO_EXT}}
LABEL = {"text": "📄 文字", "image": "🖼️ 圖片", "audio": "🎵 音頻", "video": "🎬 影片"}

SKIP_DIRS = {"windows", "program files", "program files (x86)", "programdata", "$recycle.bin",
             "system volume information", "appdata", "node_modules", ".git", "venv", ".venv",
             "__pycache__", "runtime", "site-packages", ".cache", "$windows.~bt", "recovery"}

MAX_TEXT_CHARS = 8000      # 每個文字檔最多讀幾多字
CHUNK_CHARS = 1500         # 每段長度
MAX_CHUNKS = 5             # 每個檔最多幾段
VIDEO_FRAMES = 16          # 每條影片均勻抽幾多幀
AUDIO_MAX_SEC = 300        # 音頻最多取幾多秒（模型上限約 327 秒）
MAX_FILE_MB = {"text": 50, "image": 60, "audio": 500, "video": 8000}


# ---------------- 模型 ----------------
@st.cache_resource(show_spinner="正在載入 EmbeddingGemma 2 模型（第一次約 10-30 秒）...")
def load_model():
    import torch
    from sentence_transformers import SentenceTransformer
    if torch.cuda.is_available():
        device = "cuda"
        dtype = torch.bfloat16 if torch.cuda.is_bf16_supported() else torch.float32
    else:
        device = "cpu"
        dtype = torch.float32
        try:
            torch.set_num_threads(max(1, os.cpu_count() or 1))
        except Exception:
            pass
    # 注意：模型唔支援 float16，只用 bfloat16 / float32
    model = SentenceTransformer(MODEL_DIR, device=device, model_kwargs={"torch_dtype": dtype})
    name = torch.cuda.get_device_name(0) if device == "cuda" else platform.processor() or "CPU"
    return model, device, str(dtype).replace("torch.", ""), name


def _norm(v):
    v = np.asarray(v, dtype=np.float32).reshape(-1)
    n = np.linalg.norm(v)
    return v / n if n > 0 else v


def encode_query_text(model, q):
    return _norm(model.encode(q, prompt_name="SearchQuery", normalize_embeddings=True))


def encode_doc_texts(model, texts, batch_size=8):
    e = model.encode(texts, batch_size=batch_size, normalize_embeddings=True, show_progress_bar=False)
    return np.asarray(e, dtype=np.float32)


def _encode_one(model, inp):
    return _norm(model.encode([inp], normalize_embeddings=True, show_progress_bar=False)[0])


def encode_image(model, path_or_pil):
    from PIL import Image, ImageOps
    img = path_or_pil
    if not isinstance(img, Image.Image):
        img = Image.open(img)
    img = ImageOps.exif_transpose(img).convert("RGB")
    img.thumbnail((1536, 1536))
    return _encode_one(model, {"image": img})


def ffmpeg_exe():
    try:
        import imageio_ffmpeg
        return imageio_ffmpeg.get_ffmpeg_exe()
    except Exception:
        return "ffmpeg"


def load_audio_16k(path, max_sec=AUDIO_MAX_SEC):
    """用 ffmpeg 將任何音頻/影片聲軌轉成 16kHz 單聲道 float32"""
    cmd = [ffmpeg_exe(), "-v", "quiet", "-i", str(path), "-t", str(max_sec),
           "-vn", "-f", "f32le", "-ac", "1", "-ar", "16000", "-"]
    flags = 0x08000000 if os.name == "nt" else 0  # CREATE_NO_WINDOW
    out = subprocess.run(cmd, capture_output=True, creationflags=flags).stdout
    arr = np.frombuffer(out, dtype=np.float32).copy()
    if arr.size < 1600:
        raise ValueError("讀唔到音頻內容")
    return arr


def encode_audio(model, path):
    arr = load_audio_16k(path)
    last = None
    # 唔同版本 sentence-transformers 接受嘅格式略有不同，逐個試
    for inp in ({"audio": {"array": arr, "sampling_rate": 16000}}, {"audio": arr}, {"audio": str(path)}):
        try:
            return _encode_one(model, inp)
        except Exception as ex:
            last = ex
    raise last


def video_frames(path, n=VIDEO_FRAMES):
    import cv2
    from PIL import Image
    cap = cv2.VideoCapture(str(path))
    total = int(cap.get(cv2.CAP_PROP_FRAME_COUNT) or 0)
    frames = []
    if total > 0:
        idxs = np.linspace(0, max(total - 1, 0), num=min(n, total)).astype(int)
        for i in idxs:
            cap.set(cv2.CAP_PROP_POS_FRAMES, int(i))
            ok, f = cap.read()
            if ok:
                frames.append(Image.fromarray(cv2.cvtColor(f, cv2.COLOR_BGR2RGB)))
    if not frames:  # 部分格式唔支援跳幀，逐幀讀
        cap.release(); cap = cv2.VideoCapture(str(path)); allf = []
        while len(allf) < 3000:
            ok, f = cap.read()
            if not ok: break
            allf.append(f)
        for i in np.linspace(0, len(allf) - 1, num=min(n, len(allf))).astype(int) if allf else []:
            frames.append(Image.fromarray(cv2.cvtColor(allf[i], cv2.COLOR_BGR2RGB)))
    cap.release()
    for fr in frames:
        fr.thumbnail((768, 768))
    if not frames:
        raise ValueError("讀唔到影片畫面")
    return frames


def encode_video(model, path):
    frames = video_frames(path)
    last = None
    for inp in ({"video": frames}, {"video": str(path)}, {"image": frames}):
        try:
            return _encode_one(model, inp)
        except Exception as ex:
            last = ex
    raise last


def read_text_file(path):
    ext = path.suffix.lower()
    try:
        if ext == ".pdf":
            from pypdf import PdfReader
            r = PdfReader(str(path)); out = []
            for p in r.pages[:30]:
                out.append(p.extract_text() or "")
                if sum(map(len, out)) > MAX_TEXT_CHARS: break
            return "\n".join(out)
        if ext == ".docx":
            import docx
            return "\n".join(p.text for p in docx.Document(str(path)).paragraphs)
        if ext == ".pptx":
            from pptx import Presentation
            out = []
            for s in Presentation(str(path)).slides:
                for sh in s.shapes:
                    if getattr(sh, "has_text_frame", False):
                        out.append(sh.text_frame.text)
            return "\n".join(out)
        if ext == ".xlsx":
            import openpyxl
            wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True); out = []
            for ws in wb.worksheets[:5]:
                for row in ws.iter_rows(max_row=300, values_only=True):
                    out.append(" ".join(str(c) for c in row if c is not None))
            return "\n".join(out)
        raw = path.read_bytes()[: MAX_TEXT_CHARS * 4]
        for enc in ("utf-8", "utf-16", "big5", "gb18030"):
            try:
                return raw.decode(enc)
            except Exception:
                continue
        return raw.decode("utf-8", errors="ignore")
    except Exception:
        return ""


def text_chunks(path):
    txt = " ".join(read_text_file(path).split())[:MAX_TEXT_CHARS]
    if not txt.strip():
        return []
    chunks = [txt[i:i + CHUNK_CHARS] for i in range(0, len(txt), CHUNK_CHARS)][:MAX_CHUNKS]
    title = path.stem
    return [(f"title: {title} | text: {c}", c[:300]) for c in chunks]


# ---------------- 索引 ----------------
def index_path(folder):
    h = hashlib.md5(str(Path(folder).resolve()).lower().encode("utf-8")).hexdigest()[:12]
    return INDEX_DIR / h


def load_index(folder):
    p = index_path(folder)
    if (p / "meta.json").exists() and (p / "emb.npy").exists():
        meta = json.loads((p / "meta.json").read_text(encoding="utf-8"))
        emb = np.load(p / "emb.npy")
        if len(meta["rows"]) == len(emb):
            return meta, emb
    return {"folder": str(folder), "rows": [], "files": {}}, np.zeros((0, 768), np.float32)


def save_index(folder, meta, emb):
    p = index_path(folder); p.mkdir(parents=True, exist_ok=True)
    meta["updated"] = time.strftime("%Y-%m-%d %H:%M")
    np.save(p / "emb.npy", emb.astype(np.float32))
    (p / "meta.json").write_text(json.dumps(meta, ensure_ascii=False), encoding="utf-8")


def list_indexes():
    out = []
    for d in INDEX_DIR.iterdir() if INDEX_DIR.exists() else []:
        m = d / "meta.json"
        if m.exists():
            try:
                meta = json.loads(m.read_text(encoding="utf-8"))
                out.append(meta)
            except Exception:
                pass
    return out


def scan_files(folder, mods):
    res = []
    for dirpath, dirnames, filenames in os.walk(folder):
        dirnames[:] = [d for d in dirnames if d.lower() not in SKIP_DIRS and not d.startswith(".")]
        for fn in filenames:
            p = Path(dirpath) / fn
            m = MODALITY.get(p.suffix.lower())
            if m and m in mods and not fn.startswith("~$"):
                try:
                    stt = p.stat()
                    if stt.st_size > 0 and stt.st_size <= MAX_FILE_MB[m] * 1024 * 1024:
                        res.append((str(p), m, stt.st_mtime, stt.st_size))
                except OSError:
                    pass
    return res


def build_index(model, folder, mods, progress, status):
    meta, emb = load_index(folder)
    files = scan_files(folder, mods)
    current = {f[0]: f for f in files}
    old_files = meta.get("files", {})

    # 刪除已唔存在 / 已修改的檔案
    keep = [i for i, r in enumerate(meta["rows"])
            if r["path"] in current and old_files.get(r["path"]) == [current[r["path"]][2], current[r["path"]][3]]]
    rows = [meta["rows"][i] for i in keep]
    emb = emb[keep] if len(keep) else np.zeros((0, 768), np.float32)
    done = {r["path"] for r in rows}
    files_rec = {p: v for p, v in old_files.items() if p in done}

    todo = [f for f in files if f[0] not in done]
    new_rows, new_embs, errors = [], [], []
    t0 = time.time()
    for k, (path, m, mtime, size) in enumerate(todo):
        name = Path(path).name
        status.write(f"({k + 1}/{len(todo)}) {LABEL[m]}  {name}")
        try:
            if m == "text":
                ch = text_chunks(Path(path))
                if ch:
                    e = encode_doc_texts(model, [c[0] for c in ch])
                    for (_, snip), v in zip(ch, e):
                        new_rows.append({"path": path, "mod": m, "snip": snip}); new_embs.append(_norm(v))
            elif m == "image":
                new_rows.append({"path": path, "mod": m}); new_embs.append(encode_image(model, path))
            elif m == "audio":
                new_rows.append({"path": path, "mod": m}); new_embs.append(encode_audio(model, path))
            elif m == "video":
                new_rows.append({"path": path, "mod": m}); new_embs.append(encode_video(model, path))
            files_rec[path] = [mtime, size]
        except Exception as ex:
            errors.append(f"{name}: {type(ex).__name__} {str(ex)[:120]}")
        progress.progress((k + 1) / max(len(todo), 1), text=f"已處理 {k + 1}/{len(todo)}，用時 {int(time.time() - t0)} 秒")
        # 每 50 個檔存一次，中途關閉都唔會白做
        if (k + 1) % 50 == 0 and new_embs:
            save_index(folder, {"folder": str(folder), "rows": rows + new_rows, "files": files_rec, "mods": sorted(mods)},
                       np.vstack([emb] + [np.stack(new_embs)]) if len(emb) else np.stack(new_embs))
    if new_embs:
        emb = np.vstack([emb, np.stack(new_embs)]) if len(emb) else np.stack(new_embs)
    rows = rows + new_rows
    meta = {"folder": str(folder), "rows": rows, "files": files_rec, "mods": sorted(mods)}
    save_index(folder, meta, emb)
    return len(todo), len(files), errors


# ---------------- 介面 ----------------
st.set_page_config(page_title="EmbeddingGemma 2 本地搜索", page_icon="🔎", layout="wide")
st.title("🔎 EmbeddingGemma 2 本地多模態搜索")
st.caption("文字、圖片、影片、音頻互相搜索 · 全部喺你部電腦本地運行，唔會上傳任何檔案")

try:
    model, device, dtype, devname = load_model()
except Exception as ex:
    st.error(f"模型載入失敗：{ex}\n\n請確認模型檔案喺：{MODEL_DIR}")
    st.stop()

with st.sidebar:
    st.success(f"運行裝置：{'GPU' if device == 'cuda' else 'CPU'} · {devname} · {dtype}")
    st.header("① 建立索引")
    folder = st.text_input("要搜索的資料夾路徑", placeholder=r"例如 D:\Photos 或 C:\Users\你\Documents")
    mods = st.multiselect("索引哪些類型", ["text", "image", "audio", "video"],
                          default=["text", "image", "audio", "video"], format_func=lambda x: LABEL[x])
    if folder and Path(folder).resolve().parent == Path(folder).resolve() and Path(folder).drive.upper().startswith("C"):
        st.warning("唔建議揀成個 C 盤：系統檔案太多，索引會好慢兼佔大量空間。")
    if st.button("開始建立 / 更新索引", type="primary", use_container_width=True):
        if not folder or not Path(folder).is_dir():
            st.error("資料夾唔存在，請檢查路徑。")
        elif not mods:
            st.error("請至少揀一種類型。")
        else:
            prog = st.progress(0.0, text="掃描檔案中...")
            stat = st.empty()
            n_new, n_all, errs = build_index(model, folder, set(mods), prog, stat)
            stat.empty()
            st.success(f"完成！共 {n_all} 個檔案，今次新增/更新 {n_new} 個。")
            if errs:
                with st.expander(f"{len(errs)} 個檔案略過"):
                    st.text("\n".join(errs[:300]))

    st.header("② 已建立的索引")
    idx_list = list_indexes()
    chosen = []
    for meta in idx_list:
        n_files = len(meta.get("files", {}))
        if st.checkbox(f"{meta['folder']}  ({n_files} 個檔 · {meta.get('updated', '')})", value=True,
                       key="idx_" + meta["folder"]):
            chosen.append(meta["folder"])
    if not idx_list:
        st.info("未有索引，請先喺上面建立。")

    st.header("③ 搜索設定")
    topk = st.slider("顯示結果數量", 4, 60, 12)
    show_mods = st.multiselect("只顯示哪些類型", ["text", "image", "audio", "video"],
                               default=["text", "image", "audio", "video"], format_func=lambda x: LABEL[x])


def open_in_explorer(path):
    try:
        if os.name == "nt":
            subprocess.Popen(["explorer", "/select,", os.path.normpath(path)])
        else:
            subprocess.Popen(["xdg-open", str(Path(path).parent)])
    except Exception as ex:
        st.toast(f"打開失敗：{ex}")


def run_search(qvec, exclude=None):
    all_rows, all_emb = [], []
    for f in chosen:
        meta, emb = load_index(f)
        if len(emb):
            all_rows += meta["rows"]; all_emb.append(emb)
    if not all_emb:
        st.warning("未有可用索引，請先建立索引並喺左邊勾選。")
        return
    E = np.vstack(all_emb)
    scores = E @ qvec
    best = {}
    for i in np.argsort(-scores):
        r = all_rows[i]
        if r["mod"] not in show_mods or r["path"] == exclude:
            continue
        if r["path"] not in best:
            best[r["path"]] = (float(scores[i]), r)
        if len(best) >= topk:
            break
    st.subheader(f"搜索結果（{len(best)}）")
    items = list(best.values())
    cols = st.columns(4)
    for j, (sc, r) in enumerate(items):
        p = Path(r["path"])
        with cols[j % 4]:
            with st.container(border=True):
                st.markdown(f"**{LABEL[r['mod']]}** · 相似度 `{sc:.3f}`")
                st.caption(p.name)
                try:
                    if not p.exists():
                        st.warning("檔案已被移走")
                    elif r["mod"] == "image":
                        st.image(str(p), use_container_width=True)
                    elif r["mod"] == "audio":
                        st.audio(str(p))
                    elif r["mod"] == "video":
                        st.video(str(p))
                    else:
                        st.write((r.get("snip") or "")[:220] + "…")
                except Exception:
                    st.write("（預覽失敗）")
                st.caption(str(p.parent))
                if st.button("📂 打開所在位置", key=f"open_{j}_{hash(r['path'])}"):
                    open_in_explorer(r["path"])


tabs = st.tabs(["📝 用文字搜", "🖼️ 用圖片搜", "🎵 用音頻搜", "🎬 用影片搜"])

with tabs[0]:
    q = st.text_input("輸入描述（中英文都得）", placeholder="例如：海邊日落、地盤安全帽、會議紀錄關於水庫、有狗吠聲")
    if q:
        with st.spinner("搜索中..."):
            run_search(encode_query_text(model, q))


def save_upload(up):
    suffix = Path(up.name).suffix
    tmp = Path(tempfile.gettempdir()) / f"eg2_query{suffix}"
    tmp.write_bytes(up.getbuffer())
    return tmp


with tabs[1]:
    up = st.file_uploader("上傳一張圖片，搵相似圖片 / 相關影片 / 文件", type=[e[1:] for e in IMAGE_EXT], key="up_img")
    if up:
        st.image(up, width=260)
        with st.spinner("分析圖片中..."):
            run_search(encode_image(model, save_upload(up)))

with tabs[2]:
    up = st.file_uploader("上傳音頻（mp3/wav/m4a...）", type=[e[1:] for e in AUDIO_EXT], key="up_aud")
    if up:
        st.audio(up)
        with st.spinner("分析音頻中..."):
            run_search(encode_audio(model, save_upload(up)))

with tabs[3]:
    up = st.file_uploader("上傳影片（mp4/mov...）", type=[e[1:] for e in VIDEO_EXT], key="up_vid")
    if up:
        st.video(up)
        with st.spinner("分析影片中..."):
            run_search(encode_video(model, save_upload(up)))
